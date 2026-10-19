"""Read-only course capacity snapshot; credentials via stdin, never saved."""
import argparse
import csv
import getpass
import json
import os
import sys
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from query_course import FIELDS, cell, make_url, query


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    cookie = os.environ.get('ZUEL_COOKIE') or (getpass.getpass('Cookie: ')
        if sys.stdin.isatty() else sys.stdin.readline().strip())
    if not cookie:
        raise RuntimeError('缺少 Cookie')
    source = load_workbook(args.source, read_only=True, data_only=True)
    sheet = source.active
    rows = list(sheet.values)
    headers = list(rows[1])
    module = headers.index('所属模块（适用2024版）')
    name = headers.index('课程名称')
    selected = [(i, list(row)) for i, row in enumerate(rows[2:], 3)
                if row[module] == '大类平台（经法管融通）']
    names = list(dict.fromkeys(row[name] for _, row in selected))
    args.out.mkdir(parents=True, exist_ok=False)
    records, logs, summaries = [], [], []
    for index, course in enumerate(names, 1):
        result = []
        for kind in ('public', 'cross'):
            params = SimpleNamespace(course=course, teacher='', weekday=None,
                start=None, end=None, sfym='false', sfct='false', sfxx='true',
                type=kind, szjylb='', kkdw='', skfs='')
            stamp = datetime.now().astimezone().isoformat(timespec='seconds')
            try:
                batch = query(make_url(params), cookie)
            except Exception as exc:
                logs.append([course, kind, stamp, 'ERROR', type(exc).__name__ + ': ' + str(exc)])
                (args.out / 'requests.json').write_text(json.dumps(logs, ensure_ascii=False, indent=2), encoding='utf-8')
                raise
            logs.append([course, kind, stamp, len(batch), ''])
            for row in batch:
                semester = row.get('xnxq01id')
                if semester and semester != '2026-2027-1':
                    raise RuntimeError('接口返回其他学期，停止导出: ' + str(semester))
                result.append(dict(query=course, endpoint=kind, fetched_at=stamp,
                    exact=row.get('kcmc') == course, data=row))
            if batch:
                break
        records.extend(result)
        summaries.append([course, len(result), sum(r['exact'] for r in result),
                          '有返回结果' if result else '当前会话与筛选条件下未返回'])
        (args.out / 'raw.json').write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding='utf-8')
        print(f'{index}/{len(names)} {course}: {len(result)}', flush=True)
    book = Workbook()
    book.remove(book.active)
    details = [[r['query'], r['exact'], r['endpoint'], r['fetched_at']] +
               [cell(r['data'].get(key)) for key, _ in FIELDS] for r in records]
    tables = {
        '教学班余量': ([ '查询课程', '名称精确匹配', '查询接口', '查询时间'] + [label for _, label in FIELDS], details),
        '课程汇总': (['查询课程', '教学班数', '名称精确匹配数', '状态'], summaries),
        '源表筛选记录': (['源表行号'] + headers, [[i] + row for i, row in selected]),
        '请求记录': (['查询课程', '接口', '时间', '返回条数', '错误'], logs),
        '说明': (['项目', '内容'], [
            ['源文件', str(args.source.resolve())],
            ['目标学期', '2026-2027-1；返回行含学期时逐条核验，空学期不视为已核验'],
            ['筛选条件', '所属模块（适用2024版）=大类平台（经法管融通）'],
            ['接口策略', '公选接口有结果则保留；无结果时跨专业补查'],
            ['范围', 'sfym=false, sfct=false, sfxx=true；当前账号可见结果，非全校完整开课清单'],
            ['余量', '原样保留接口 syrs 字段，空值不视为零；数据为查询时快照'],
            ['匹配', '接口可能模糊匹配，请结合返回课程编号及名称精确匹配标记判断'],
        ])}
    for title, (cols, values) in tables.items():
        ws = book.create_sheet(title)
        for row in [cols] + values:
            ws.append(row)
            for c in ws[ws.max_row]:
                if c.data_type == 'f':
                    c.data_type = 's'
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    book.save(args.out / '经法管融通课程余量.xlsx')
    with (args.out / '教学班余量.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(tables['教学班余量'][0])
        writer.writerows(details)
    (args.out / 'requests.json').write_text(json.dumps(logs, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(dict(source_rows=len(selected), courses=len(names), classes=len(records),
        courses_found=sum(bool(r[1]) for r in summaries)), ensure_ascii=False))


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
