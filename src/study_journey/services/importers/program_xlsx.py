"""培养方案 Excel → data/program.json。

只导出培养方案本身（课程、学分、开课学期、先修条件），
不导出表格里的个人字段（完成学期、分数、意愿度）。

用法：uv run --group importers python -m study_journey.services.importers.program_xlsx \
        src_data/培养方案.xlsx src/study_journey/data/program.json
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

CATEGORIES = {
    "通识教育": "general_education",
    "大类平台": "major_platform",
    "专业方向": "professional_education",
    "专业教育": "professional_education",
    "素质教育": "quality_education",
    "实践教育": "practice_education",
}
CATEGORY_NAMES = {
    "general_education": "通识教育",
    "major_platform": "大类平台",
    "professional_education": "专业教育",
    "quality_education": "素质教育",
    "practice_education": "实践教育",
}
# 毕业要求（docs/16 §2.1 定稿）
REQUIREMENTS = [
    ("general_education", "required", "必修", 32),
    ("general_education", "elective", "选修", 16),
    ("major_platform", "required", "必修", 30),
    ("major_platform", "elective", "选修", 10),
    ("major_platform", "interdisciplinary", "经法管融通", 6),
    ("professional_education", "required", "必修", 18),
    ("professional_education", "elective", "选修", 20),
    ("quality_education", "required", "必修", 6),
    ("quality_education", "elective", "选修", 6),
    ("practice_education", "required", "必修", 16),
]
CODE = re.compile(r"^[A-Z]\d{7}$")


def text(value) -> str | None:
    if value is None:
        return None
    s = re.sub(r"\s+", "", str(value).replace("\xa0", ""))
    return s or None


def number(value) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def nature_of(value: str | None) -> str | None:
    """'必修\\n32' → '必修'。"""
    s = text(value)
    if not s:
        return None
    for n in ("必修", "选修"):
        if s.startswith(n):
            return n
    return None


PREREQ_ALIASES = {"高等数学(上)": "高等数学(1)"}


def first_term(course: dict) -> int:
    m = re.match(r"\d", course["offered_term"] or "")
    return int(m.group()) if m else 9


def resolve_prerequisites(courses: list[dict]) -> None:
    """表格里的先修条件写的是课程名称（且不完全规范），尽量对应到课程号；对不上就留空。"""
    for c in courses:
        raw = c.pop("prerequisite")
        name = PREREQ_ALIASES.get(raw, raw)
        code = None
        if name and not name.isdigit():
            exact = [o for o in courses if o["name"] == name]
            prefix = sorted((o for o in courses if o["name"].startswith(name)), key=first_term)
            match = (exact or prefix or [None])[0]
            code = match["code"] if match and match is not c else None
        c["prerequisite_code"] = code


def parse(path: Path) -> dict:
    wb = load_workbook(path, data_only=True)
    courses: list[dict] = []
    for ws in wb.worksheets:
        sheet_category = next((c for k, c in CATEGORIES.items() if ws.title.startswith(k)), None)
        rows = list(ws.iter_rows(values_only=True))
        header = [text(h) or "" for h in rows[0]]
        has_segment = header[0] == "培养环节"
        category, nature = sheet_category, None
        for row_no, row in enumerate(rows[1:], 2):
            cells = list(row)
            if has_segment:
                seg = text(cells[0])
                if seg == "培养环节":  # 第二张子表的表头
                    continue
                if seg:
                    category = next((c for k, c in CATEGORIES.items() if seg.startswith(k)), category)
                cells = cells[1:]
            first = text(cells[0])
            if first and first.startswith(("选修类别", "跨专业类别")):
                break  # 下面是个人已选课程，不属于培养方案
            nature = nature_of(cells[0]) or nature
            code = text(cells[1])
            if not code or not CODE.match(code) or not nature or not category:
                continue
            lecture, experiment, extra = (int(number(cells[i])) for i in (5, 6, 7))
            if code == "B2490014":  # 形势与政策(1—8)：学校按学期拆成 8 门，每门 0.25 学分
                for n in range(1, 9):
                    courses.append({
                        "code": f"B24900{n}4", "name": f"形势与政策({n})",
                        "category": category, "nature": nature, "offered_term": str(n),
                        "credits": number(cells[4]) / 8, "lecture_hours": lecture // 8,
                        "experiment_hours": 0, "extra_hours": 0, "prerequisite": None,
                        "source": f"{ws.title}!{row_no}",
                    })
                continue
            courses.append({
                "code": code,
                "name": text(cells[2]),
                "category": category,
                "nature": nature,
                "offered_term": text(cells[3]),
                "credits": number(cells[4]),
                "lecture_hours": lecture,
                "experiment_hours": experiment,
                "extra_hours": extra,
                "prerequisite": text(cells[8]),
                "source": f"{ws.title}!{row_no}",
            })
    resolve_prerequisites(courses)
    return {
        "program": "金融学（数字金融）2024 级",
        "categories": CATEGORY_NAMES,
        "requirements": [
            {"category": c, "type": t, "name": n, "credits": cr} for c, t, n, cr in REQUIREMENTS
        ],
        "courses": courses,
    }


def main(src: str, out: str) -> None:
    data = parse(Path(src))
    Path(out).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"courses: {len(data['courses'])} -> {out}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
